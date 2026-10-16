"""
Exportador corporativo para planilhas Excel (.xlsx) e relatórios PDF Paisagem (.pdf).
Aplica regras financeiras cumulativas de formatação e proteção contra PermissionError.
"""
import os
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.data_processor import (
    formatar_moeda_float,
    formatar_float_moeda,
    calcular_total_a_pagar_veiculo,
    obter_estilo_financeiro
)
from core.logger_config import obter_logger

logger = obter_logger("RelatorioExporter")

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.graphics.shapes import Drawing, Wedge, Circle, Rect, String, Line
    REPORTLAB_DISPONIVEL = True
except ImportError:
    REPORTLAB_DISPONIVEL = False


if REPORTLAB_DISPONIVEL:
    class NumberedCanvas(canvas.Canvas):
        """Canvas dinâmico de 2 passos para numeração de páginas 'Página X de Y'."""
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_paginas = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_number(num_paginas)
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

        def draw_page_number(self, total_paginas):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            largura, altura = landscape(A4)
            data_atual = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
            texto_esq = f"Relatório Gerado em: {data_atual} | Sistema de Gestão de Frotas (CE)"
            self.drawString(36, 20, texto_esq)
            texto_dir = f"Página {self._pageNumber} de {total_paginas}"
            self.drawRightString(largura - 36, 20, texto_dir)
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 32, largura - 36, 32)
            self.restoreState()
else:
    NumberedCanvas = None


class FrotaRelatorioExporter:
    """Exportador para planilhas Excel formatadas e relatórios PDF em layout Paisagem."""

    @staticmethod
    def exportar_excel(df: pd.DataFrame, caminho_arquivo: str) -> str:
        """
        Gera e formata a planilha Excel de saída com estilo corporativo.
        Trata PermissionError salvando com timestamp para evitar perda de dados.
        Retorna o caminho efetivo do arquivo gravado.
        """
        caminho_efetivo = caminho_arquivo
        pasta_dir = os.path.dirname(os.path.abspath(caminho_efetivo))
        if pasta_dir:
            os.makedirs(pasta_dir, exist_ok=True)
        try:
            df.to_excel(caminho_efetivo, index=False, engine="openpyxl")
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            pasta, nome = os.path.split(caminho_arquivo)
            raiz, ext = os.path.splitext(nome)
            caminho_efetivo = os.path.join(pasta, f"{raiz}_{timestamp}{ext}")
            df.to_excel(caminho_efetivo, index=False, engine="openpyxl")

        wb = load_workbook(caminho_efetivo)
        ws = wb.active

        cor_cabecalho = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fonte_linhas = Font(name="Calibri", size=10)
        alinhamento_centro = Alignment(horizontal="center", vertical="center")
        alinhamento_direita = Alignment(horizontal="right", vertical="center")
        alinhamento_esquerda = Alignment(horizontal="left", vertical="center")
        borda_fina = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for cell in ws[1]:
            cell.fill = cor_cabecalho
            cell.font = fonte_cabecalho
            cell.alignment = alinhamento_centro

        cor_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        col_lic_idx = None
        col_ipva_idx = None
        col_multas_idx = None
        col_total_idx = None
        col_status_idx = None

        for col_idx, col_name in enumerate(df.columns, start=1):
            cn = str(col_name).strip().upper()
            if "LICENCIAMENTO" in cn:
                col_lic_idx = col_idx
            elif "IPVA" in cn:
                col_ipva_idx = col_idx
            elif "MULTA" in cn:
                col_multas_idx = col_idx
            elif "TOTAL" in cn:
                col_total_idx = col_idx
            elif "STATUS" in cn:
                col_status_idx = col_idx

        for row_idx in range(2, ws.max_row + 1):
            is_zebra = (row_idx % 2 == 0)
            fill_padrao = cor_zebra if is_zebra else None

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if fill_padrao:
                    cell.fill = fill_padrao
                cell.font = fonte_linhas
                cell.border = borda_fina
                val_txt = str(cell.value or "")

                is_col_financeira = col_idx in [col_lic_idx, col_ipva_idx, col_multas_idx, col_total_idx]
                val_txt_upper = val_txt.upper()

                if is_col_financeira:
                    if val_txt_upper == "PAGO":
                        cell.alignment = alinhamento_centro
                        # Fundo verde pastel suave e fonte verde escuro elegante
                        cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                    elif "RESTRIÇÃO" in val_txt_upper or "RESTRICAO" in val_txt_upper:
                        cell.alignment = alinhamento_centro
                        # Fundo laranja pastel suave e fonte laranja escuro elegante
                        cell.fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
                        cell.font = Font(name="Calibri", size=10, bold=True, color="9A3412")
                    else:
                        cell.alignment = alinhamento_direita
                        val_float = formatar_moeda_float(val_txt)
                        is_lic = (col_idx == col_lic_idx)
                        is_tot = (col_idx == col_total_idx)

                        bold, underline, cor_hex = obter_estilo_financeiro(val_float, is_licenciamento=is_lic)
                        if is_tot:
                            bold = True
                            if val_float > 0:
                                cor_hex = "#15803D"

                        sublinhado_estilo = 'single' if underline else None
                        cor_fonte = cor_hex.replace("#", "") if cor_hex else '000000'

                        cell.font = Font(
                            name="Calibri",
                            size=10,
                            bold=bold,
                            underline=sublinhado_estilo,
                            color=cor_fonte
                        )
                elif col_idx == col_status_idx:
                    cell.alignment = alinhamento_centro
                    if "RESTRIÇÃO" in val_txt_upper or "RESTRICAO" in val_txt_upper:
                        cell.fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
                        cell.font = Font(name="Calibri", size=10, bold=True, color="9A3412")
                    elif "SEM DÉBITO" in val_txt_upper:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                    elif "COM DÉBITO" in val_txt_upper:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                    elif "FALHA" in val_txt_upper:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="B45309")
                elif "PLACA" in str(ws.cell(row=1, column=col_idx).value or "").upper():
                    cell.alignment = alinhamento_centro
                    cell.font = Font(name="Calibri", size=10, bold=True)
                elif "RENAVAM" in str(ws.cell(row=1, column=col_idx).value or "").upper():
                    cell.alignment = alinhamento_centro
                else:
                    cell.alignment = alinhamento_esquerda

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        ws.freeze_panes = "A2"
        wb.save(caminho_efetivo)
        logger.info(f"Relatório Excel gravado com sucesso: {caminho_efetivo}")
        return caminho_efetivo

    @staticmethod
    def exportar_pdf(
        df: pd.DataFrame,
        caminho_arquivo: str,
        totais: dict,
        filtros_ativos: list[str],
        identificador_frota: str = ""
    ) -> str:
        """
        Gera relatório executivo em PDF no formato Paisagem (A4) com ReportLab.
        Aplica regras financeiras cumulativas de estilo e proteção de sobreescrita.
        """
        if not REPORTLAB_DISPONIVEL:
            raise RuntimeError("ReportLab não está instalado no ambiente.")

        caminho_efetivo = caminho_arquivo
        pasta_dir = os.path.dirname(os.path.abspath(caminho_efetivo))
        if pasta_dir:
            os.makedirs(pasta_dir, exist_ok=True)

        try:
            with open(caminho_efetivo, "ab"):
                pass
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            pasta, nome = os.path.split(caminho_arquivo)
            raiz, ext = os.path.splitext(nome)
            caminho_efetivo = os.path.join(pasta, f"{raiz}_{timestamp}{ext}")

        doc = SimpleDocTemplate(
            caminho_efetivo,
            pagesize=landscape(A4),
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        estilo_titulo = ParagraphStyle('TituloDoc', parent=styles['Heading1'], fontSize=15, leading=18, textColor=colors.HexColor("#1E3A8A"), fontName="Helvetica-Bold", spaceAfter=2)
        estilo_sub = ParagraphStyle('SubDoc', parent=styles['Normal'], fontSize=8.5, leading=11, textColor=colors.HexColor("#64748B"), spaceAfter=8)
        estilo_th = ParagraphStyle('Th', parent=styles['Normal'], fontSize=7.5, leading=9.5, fontName="Helvetica-Bold", textColor=colors.white, alignment=1)
        estilo_td = ParagraphStyle('Td', parent=styles['Normal'], fontSize=7, leading=9, textColor=colors.HexColor("#1F2937"), alignment=1)
        estilo_td_left = ParagraphStyle('TdLeft', parent=styles['Normal'], fontSize=7, leading=9, textColor=colors.HexColor("#1F2937"), alignment=0)
        estilo_td_right = ParagraphStyle('TdRight', parent=styles['Normal'], fontSize=7, leading=9, textColor=colors.HexColor("#1F2937"), alignment=2)

        elementos = []
        elementos.append(Paragraph("RELATÓRIO CONSOLIDADO DE DÉBITOS VEICULARES", estilo_titulo))
        sub_texto = f"Portais Consultados: Detran-CE & Sefaz-CE | Filtros Ativos: {', '.join(filtros_ativos)} | Total de Veículos: {len(df)}"
        if identificador_frota:
            sub_texto += f" | Frota: {identificador_frota}"
        elementos.append(Paragraph(sub_texto, estilo_sub))

        # CARDS DE TOTAIS
        dados_cards = [[
            Paragraph(f"<b>VEÍCULOS</b><br/><font size=10 color='#1E3A8A'><b>{len(df)}</b></font>", estilo_td),
            Paragraph(f"<b>LICENCIAMENTO</b><br/><font size=10 color='#B45309'><b>{totais.get('lic', 'R$ 0,00')}</b></font>", estilo_td),
            Paragraph(f"<b>IPVA EM ABERTO</b><br/><font size=10 color='#B91C1C'><b>{totais.get('ipva', 'R$ 0,00')}</b></font>", estilo_td),
            Paragraph(f"<b>MULTAS DETRAN</b><br/><font size=10 color='#7C2D12'><b>{totais.get('multas', 'R$ 0,00')}</b></font>", estilo_td),
            Paragraph(f"<b>TOTAL A PAGAR</b><br/><font size=10 color='#15803D'><b>{totais.get('total', 'R$ 0,00')}</b></font>", estilo_td)
        ]]
        tabela_cards = Table(dados_cards, colWidths=[150]*5)
        tabela_cards.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elementos.append(tabela_cards)
        elementos.append(Spacer(1, 10))

        # TABELA DE DADOS
        colunas_df = list(df.columns)
        cabecalho = [Paragraph(str(col), estilo_th) for col in colunas_df]
        tabela_linhas = [cabecalho]

        largura_total_util = 770
        largura_por_col = largura_total_util / max(len(colunas_df), 1)
        larguras = [largura_por_col] * len(colunas_df)

        for _, row in df.iterrows():
            linha_celulas = []
            for col_nome in colunas_df:
                val = str(row[col_nome] if pd.notna(row[col_nome]) else "")
                val_upper = val.upper()
                col_upper = str(col_nome).upper()

                if "STATUS" in col_upper:
                    if "RESTRIÇÃO" in val_upper or "RESTRICAO" in val_upper:
                        txt_p = f"<font color='#C2410C'><b>{val}</b></font>"
                    elif "SEM DÉBITO" in val_upper:
                        txt_p = f"<font color='#166534'><b>{val}</b></font>"
                    elif "COM DÉBITO" in val_upper:
                        txt_p = f"<font color='#991B1B'><b>{val}</b></font>"
                    elif "FALHA" in val_upper:
                        txt_p = f"<font color='#B45309'><b>{val}</b></font>"
                    else:
                        txt_p = val
                    linha_celulas.append(Paragraph(txt_p, estilo_td))
                elif any(c in col_upper for c in ["LICENCIAMENTO", "IPVA", "MULTA", "TOTAL"]):
                    if val_upper == "PAGO":
                        txt_p = "<font color='#166534'><b>PAGO</b></font>"
                        linha_celulas.append(Paragraph(txt_p, estilo_td))
                    elif "RESTRIÇÃO" in val_upper or "RESTRICAO" in val_upper:
                        txt_p = "<font color='#C2410C'><b>Restrição</b></font>"
                        linha_celulas.append(Paragraph(txt_p, estilo_td))
                    else:
                        val_float = formatar_moeda_float(val)
                        is_lic = "LICENCIAMENTO" in col_upper
                        is_tot = "TOTAL" in col_upper
                        bold, underline, cor_hex = obter_estilo_financeiro(val_float, is_licenciamento=is_lic)
                        if is_tot:
                            bold = True
                            if val_float > 0:
                                cor_hex = "#15803D"

                        conteudo = val
                        if underline:
                            conteudo = f"<u>{conteudo}</u>"
                        if bold:
                            conteudo = f"<b>{conteudo}</b>"
                        if cor_hex:
                            conteudo = f"<font color='{cor_hex}'>{conteudo}</font>"

                        linha_celulas.append(Paragraph(conteudo, estilo_td_right))
                elif "PLACA" in col_upper:
                    linha_celulas.append(Paragraph(f"<b>{val}</b>", estilo_td))
                else:
                    linha_celulas.append(Paragraph(val, estilo_td_left))
            tabela_linhas.append(linha_celulas)

        tabela_dados = Table(tabela_linhas, colWidths=larguras, repeatRows=1)
        tabela_dados.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1F497D")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0,0), (-1,-1), 0.3, colors.HexColor("#E2E8F0")),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        elementos.append(tabela_dados)

        doc.build(elementos, canvasmaker=NumberedCanvas)
        logger.info(f"Relatório PDF gerado com sucesso: {caminho_efetivo}")
        return caminho_efetivo

    @staticmethod
    def exportar_pdf_bi(
        df: pd.DataFrame,
        caminho_arquivo: str,
        totais: dict | None = None,
        filtros_ativos: list[str] | None = None,
        identificador_frota: str = ""
    ) -> str:
        """
        Gera o 2º relatório executivo em PDF (Paisagem A4) contendo o Painel Analítico de BI:
        - Gráfico Vetorial de Rosca (Donut Chart) de Composição Financeira por Tributo
        - Medidor Vetorial de Regularidade & Saúde da Frota + Indicadores Operacionais
        - Gráfico Vetorial Horizontal do Top 10 Veículos de Maior Impacto
        - Tabela Analítica Detalhada do Top 10 Veículos Críticos
        - Tabela Consolidada por Proprietário / Centro de Custo (quando presente na planilha)
        """
        if not REPORTLAB_DISPONIVEL:
            raise RuntimeError("ReportLab não está instalado no ambiente.")

        caminho_efetivo = caminho_arquivo
        pasta_dir = os.path.dirname(os.path.abspath(caminho_efetivo))
        if pasta_dir:
            os.makedirs(pasta_dir, exist_ok=True)

        try:
            with open(caminho_efetivo, "ab"):
                pass
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            pasta, nome = os.path.split(caminho_arquivo)
            raiz, ext = os.path.splitext(nome)
            caminho_efetivo = os.path.join(pasta, f"{raiz}_{timestamp}{ext}")

        filtros_ativos = filtros_ativos or ["IPVA", "Licenciamento", "Multas"]
        totais = totais or {}

        # Localizar colunas na planilha consolidada
        col_placa = next((c for c in df.columns if "PLACA" in str(c).upper()), df.columns[0])
        col_renavam = next((c for c in df.columns if any(k in str(c).upper() for k in ["RENAVAM", "RENAVAN", "RNV"])), df.columns[min(1, len(df.columns) - 1)])
        col_prop = next((c for c in df.columns if any(k in str(c).upper() for k in ["PROPRIET", "CLIENTE"])), None)
        col_lic = next((c for c in df.columns if "LICENCIAMENTO" in str(c).upper()), None)
        col_ipva = next((c for c in df.columns if "IPVA" in str(c).upper()), None)
        col_multas = next((c for c in df.columns if "MULTA" in str(c).upper()), None)
        col_total = next((c for c in df.columns if "TOTAL" in str(c).upper()), None)
        col_status = next((c for c in df.columns if "STATUS" in str(c).upper()), None)

        total_veiculos = len(df)
        soma_lic = 0.0
        soma_ipva = 0.0
        soma_multas = 0.0
        soma_total = 0.0

        qtd_pagos = 0
        qtd_debitos = 0
        qtd_restricao = 0
        qtd_falhas = 0
        qtd_boletos = 0

        ranking_veiculos = []
        agrupamento_prop: dict[str, dict] = {}

        for _, row in df.iterrows():
            placa = str(row[col_placa] if pd.notna(row[col_placa]) else "").strip().upper()
            renavam = str(row[col_renavam] if pd.notna(row[col_renavam]) else "").strip()
            prop = str(row[col_prop] if (col_prop and pd.notna(row[col_prop])) else "").strip() or "NÃO INFORMADO"

            lic_s = str(row[col_lic] if (col_lic and pd.notna(row[col_lic])) else "-")
            ipva_s = str(row[col_ipva] if (col_ipva and pd.notna(row[col_ipva])) else "-")
            multas_s = str(row[col_multas] if (col_multas and pd.notna(row[col_multas])) else "-")
            status_s = str(row[col_status] if (col_status and pd.notna(row[col_status])) else "")

            v_lic = formatar_moeda_float(lic_s)
            v_ipva = formatar_moeda_float(ipva_s)
            v_mul = formatar_moeda_float(multas_s)
            v_tot = (
                formatar_moeda_float(str(row[col_total]))
                if (col_total and pd.notna(row[col_total]) and str(row[col_total]).strip() not in ["-", ""])
                else calcular_total_a_pagar_veiculo(lic_s, ipva_s, multas_s)
            )

            soma_lic += v_lic
            soma_ipva += v_ipva
            soma_multas += v_mul
            soma_total += v_tot

            if v_lic > 0 or v_ipva > 0:
                qtd_boletos += 1

            textos_concat = f"{lic_s} {ipva_s} {multas_s} {status_s}".lower()
            tem_falha = any(t in textos_concat for t in ["falha", "erro", "timeout", "inválid", "invalid", "não localizado", "nao localizado"])
            nao_cons = all(str(v).strip() in ["-", "N/A", "Não Consultado", ""] for v in [lic_s, ipva_s, multas_s])
            is_restr = "restriç" in textos_concat or "restric" in textos_concat

            if is_restr:
                qtd_restricao += 1

            if tem_falha or nao_cons:
                qtd_falhas += 1
                sit_label = "Falha / Pendente"
            elif v_tot > 0:
                qtd_debitos += 1
                sit_label = "Com Débitos / Restrição" if is_restr else "Com Débitos Pendentes"
                ranking_veiculos.append({
                    "placa": placa,
                    "renavam": renavam,
                    "prop": prop if col_prop else "",
                    "lic_s": lic_s,
                    "ipva_s": ipva_s,
                    "multas_s": multas_s,
                    "total_f": v_tot,
                    "total_s": formatar_float_moeda(v_tot),
                    "status": sit_label
                })
            else:
                qtd_pagos += 1
                sit_label = "Restrição" if is_restr else "Regular / Quitado"

            if col_prop:
                grp = agrupamento_prop.setdefault(prop, {
                    "qtd_total": 0,
                    "qtd_debitos": 0,
                    "lic": 0.0,
                    "ipva": 0.0,
                    "multas": 0.0,
                    "total": 0.0
                })
                grp["qtd_total"] += 1
                if v_tot > 0 and not tem_falha:
                    grp["qtd_debitos"] += 1
                grp["lic"] += v_lic
                grp["ipva"] += v_ipva
                grp["multas"] += v_mul
                grp["total"] += v_tot

        ranking_veiculos.sort(key=lambda x: x["total_f"], reverse=True)
        top10 = ranking_veiculos[:10]

        soma_bruta = soma_lic + soma_ipva + soma_multas
        abatimento = max(0.0, soma_bruta - soma_total)
        auditados = qtd_pagos + qtd_debitos + qtd_falhas

        doc = SimpleDocTemplate(
            caminho_efetivo,
            pagesize=landscape(A4),
            leftMargin=36,
            rightMargin=36,
            topMargin=32,
            bottomMargin=42
        )

        styles = getSampleStyleSheet()
        estilo_titulo = ParagraphStyle(
            'TituloBI', parent=styles['Heading1'],
            fontSize=15, leading=18, textColor=colors.HexColor("#0F172A"),
            fontName="Helvetica-Bold", spaceAfter=2
        )
        estilo_sub = ParagraphStyle(
            'SubBI', parent=styles['Normal'],
            fontSize=8.5, leading=11, textColor=colors.HexColor("#475569"),
            spaceAfter=8
        )
        estilo_sec = ParagraphStyle(
            'SecaoBI', parent=styles['Heading2'],
            fontSize=10.5, leading=13, textColor=colors.HexColor("#1E3A8A"),
            fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=5
        )
        estilo_th = ParagraphStyle(
            'ThBI', parent=styles['Normal'],
            fontSize=7.5, leading=9.5, fontName="Helvetica-Bold",
            textColor=colors.white, alignment=1
        )
        estilo_td = ParagraphStyle(
            'TdBI', parent=styles['Normal'],
            fontSize=7.2, leading=9.2, textColor=colors.HexColor("#1E293B"), alignment=1
        )
        estilo_td_left = ParagraphStyle(
            'TdLeftBI', parent=styles['Normal'],
            fontSize=7.2, leading=9.2, textColor=colors.HexColor("#1E293B"), alignment=0
        )
        estilo_td_right = ParagraphStyle(
            'TdRightBI', parent=styles['Normal'],
            fontSize=7.2, leading=9.2, textColor=colors.HexColor("#1E293B"), alignment=2
        )

        elementos = []
        elementos.append(Paragraph("PAINEL ANALÍTICO DE BI & INTELIGÊNCIA DE FROTA (CE)", estilo_titulo))
        sub_texto = (
            f"Portais Auditados: Detran-CE & Sefaz-CE | Motores Ativos: {', '.join(filtros_ativos)} | "
            f"Total da Frota: {total_veiculos} veículos"
        )
        if identificador_frota:
            sub_texto += f" | Frota: {identificador_frota}"
        elementos.append(Paragraph(sub_texto, estilo_sub))

        # 1. FAIXA DE KPI CARDS EXECUTIVOS
        p_lic = (soma_lic / soma_bruta * 100.0) if soma_bruta > 0 else 0.0
        p_ipva = (soma_ipva / soma_bruta * 100.0) if soma_bruta > 0 else 0.0
        p_mul = (soma_multas / soma_bruta * 100.0) if soma_bruta > 0 else 0.0

        dados_kpis = [[
            Paragraph(
                f"<b>FROTA AUDITADA</b><br/><font size=10 color='#1E3A8A'><b>{auditados} / {total_veiculos}</b></font>"
                f"<br/><font size=6.5 color='#64748B'>{qtd_pagos} regulares • {qtd_debitos} c/ débito</font>",
                estilo_td
            ),
            Paragraph(
                f"<b>LICENCIAMENTO</b><br/><font size=10 color='#D97706'><b>{formatar_float_moeda(soma_lic)}</b></font>"
                f"<br/><font size=6.5 color='#64748B'>Detran-CE ({p_lic:.0f}% do bruto)</font>",
                estilo_td
            ),
            Paragraph(
                f"<b>IPVA (SEFAZ-CE)</b><br/><font size=10 color='#E11D48'><b>{formatar_float_moeda(soma_ipva)}</b></font>"
                f"<br/><font size=6.5 color='#64748B'>Sefaz-CE ({p_ipva:.0f}% do bruto)</font>",
                estilo_td
            ),
            Paragraph(
                f"<b>MULTAS DETRAN</b><br/><font size=10 color='#EA580C'><b>{formatar_float_moeda(soma_multas)}</b></font>"
                f"<br/><font size=6.5 color='#64748B'>Infrações ({p_mul:.0f}% do bruto)</font>",
                estilo_td
            ),
            Paragraph(
                f"<b>TOTAL LÍQUIDO A PAGAR</b><br/><font size=10 color='#15803D'><b>{formatar_float_moeda(soma_total)}</b></font>"
                f"<br/><font size=6.5 color='#64748B'>Sem duplicidade ({qtd_debitos} veíc.)</font>",
                estilo_td
            ),
        ]]
        tabela_kpis = Table(dados_kpis, colWidths=[154] * 5)
        tabela_kpis.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela_kpis)
        elementos.append(Spacer(1, 8))

        # 2. PAINEL GRÁFICO VETORIAL (3 CARDS LADO A LADO: DONUT + SAÚDE + TOP 10)
        painel_draw = FrotaRelatorioExporter._construir_painel_grafico_bi(
            total_veiculos=total_veiculos,
            auditados=auditados,
            soma_lic=soma_lic,
            soma_ipva=soma_ipva,
            soma_multas=soma_multas,
            soma_total=soma_total,
            qtd_pagos=qtd_pagos,
            qtd_debitos=qtd_debitos,
            qtd_restricao=qtd_restricao,
            qtd_falhas=qtd_falhas,
            qtd_boletos=qtd_boletos,
            top10=top10
        )
        elementos.append(painel_draw)
        elementos.append(Spacer(1, 10))

        # 3. TABELA ANALÍTICA: TOP 10 VEÍCULOS DE MAIOR IMPACTO FINANCEIRO
        elementos.append(Paragraph("RANKING DETALHADO — TOP 10 VEÍCULOS DE MAIOR IMPACTO FINANCEIRO", estilo_sec))
        if not top10:
            elementos.append(Paragraph("Nenhum veículo com débito pendente foi registrado nesta auditoria.", estilo_sub))
        else:
            cab_top10 = [
                Paragraph("#", estilo_th),
                Paragraph("Placa", estilo_th),
                Paragraph("RENAVAM", estilo_th),
                Paragraph("Proprietário / Cliente", estilo_th),
                Paragraph("Licenciamento", estilo_th),
                Paragraph("IPVA", estilo_th),
                Paragraph("Multas", estilo_th),
                Paragraph("Total Líquido", estilo_th),
                Paragraph("% Passivo", estilo_th),
                Paragraph("Situação", estilo_th),
            ]
            linhas_top10 = [cab_top10]
            for rank_idx, item in enumerate(top10, start=1):
                pct_item = (item["total_f"] / soma_total * 100.0) if soma_total > 0 else 0.0
                cor_rank = "#DC2626" if rank_idx == 1 else ("#EA580C" if rank_idx in (2, 3) else "#1E3A8A")
                linhas_top10.append([
                    Paragraph(f"<font color='{cor_rank}'><b>{rank_idx}º</b></font>", estilo_td),
                    Paragraph(f"<b>{item['placa']}</b>", estilo_td),
                    Paragraph(item["renavam"], estilo_td),
                    Paragraph((item["prop"] or "-")[:32], estilo_td_left),
                    Paragraph(item["lic_s"], estilo_td_right),
                    Paragraph(item["ipva_s"], estilo_td_right),
                    Paragraph(item["multas_s"], estilo_td_right),
                    Paragraph(f"<font color='#15803D'><b>{item['total_s']}</b></font>", estilo_td_right),
                    Paragraph(f"<b>{pct_item:.1f}%</b>", estilo_td),
                    Paragraph(f"<font color='#991B1B'><b>{item['status']}</b></font>", estilo_td),
                ])

            larg_top10 = [28, 56, 72, 148, 76, 72, 76, 84, 52, 106]
            tab_top10 = Table(linhas_top10, colWidths=larg_top10, repeatRows=1)
            tab_top10.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ('INNERGRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            elementos.append(tab_top10)

        # 4. TABELA CONSOLIDADA POR PROPRIETÁRIO / CENTRO DE CUSTO (QUANDO DISPONÍVEL)
        if col_prop and agrupamento_prop:
            elementos.append(Spacer(1, 8))
            elementos.append(Paragraph("CONSOLIDAÇÃO FINANCEIRA POR PROPRIETÁRIO / CENTRO DE CUSTO", estilo_sec))
            cab_prop = [
                Paragraph("Proprietário / Cliente", estilo_th),
                Paragraph("Veículos", estilo_th),
                Paragraph("Com Débito", estilo_th),
                Paragraph("Licenciamento", estilo_th),
                Paragraph("IPVA", estilo_th),
                Paragraph("Multas", estilo_th),
                Paragraph("Total Líquido a Pagar", estilo_th),
                Paragraph("% Participação", estilo_th),
            ]
            linhas_prop = [cab_prop]
            lista_props = sorted(agrupamento_prop.items(), key=lambda kv: kv[1]["total"], reverse=True)
            for nome_p, d_p in lista_props:
                pct_p = (d_p["total"] / soma_total * 100.0) if soma_total > 0 else 0.0
                linhas_prop.append([
                    Paragraph(f"<b>{nome_p[:38]}</b>", estilo_td_left),
                    Paragraph(str(d_p["qtd_total"]), estilo_td),
                    Paragraph(str(d_p["qtd_debitos"]), estilo_td),
                    Paragraph(formatar_float_moeda(d_p["lic"]), estilo_td_right),
                    Paragraph(formatar_float_moeda(d_p["ipva"]), estilo_td_right),
                    Paragraph(formatar_float_moeda(d_p["multas"]), estilo_td_right),
                    Paragraph(f"<font color='#15803D'><b>{formatar_float_moeda(d_p['total'])}</b></font>", estilo_td_right),
                    Paragraph(f"<b>{pct_p:.1f}%</b>", estilo_td),
                ])

            larg_prop = [210, 62, 68, 88, 88, 88, 98, 68]
            tab_prop = Table(linhas_prop, colWidths=larg_prop, repeatRows=1)
            tab_prop.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ('INNERGRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            elementos.append(tab_prop)

        doc.build(elementos, canvasmaker=NumberedCanvas)
        logger.info(f"Relatório PDF de BI gerado com sucesso: {caminho_efetivo}")
        return caminho_efetivo

    @staticmethod
    def _construir_painel_grafico_bi(
        total_veiculos: int,
        auditados: int,
        soma_lic: float,
        soma_ipva: float,
        soma_multas: float,
        soma_total: float,
        qtd_pagos: int,
        qtd_debitos: int,
        qtd_restricao: int,
        qtd_falhas: int,
        qtd_boletos: int,
        top10: list[dict]
    ) -> "Drawing":
        """Constrói o Drawing vetorial de 770x226pt com os 3 cards do Painel Analítico de BI."""
        d = Drawing(770, 226)
        soma_bruta = soma_lic + soma_ipva + soma_multas
        abatimento = max(0.0, soma_bruta - soma_total)

        c_bg_card = colors.HexColor("#F8FAFC")
        c_border = colors.HexColor("#CBD5E1")
        c_title = colors.HexColor("#0F172A")
        c_sub = colors.HexColor("#475569")
        c_muted = colors.HexColor("#64748B")
        c_track = colors.HexColor("#E2E8F0")

        # ======================================================================
        # CARD 1 (x=0..250): COMPOSIÇÃO FINANCEIRA POR TRIBUTO (DONUT CHART)
        # ======================================================================
        d.add(Rect(0, 0, 250, 226, rx=8, ry=8, fillColor=c_bg_card, strokeColor=c_border, strokeWidth=0.8))
        d.add(Rect(12, 223, 226, 3, fillColor=colors.HexColor("#F59E0B"), strokeColor=None))
        d.add(String(12, 207, "Composição Financeira por Tributo", fontName="Helvetica-Bold", fontSize=8.5, fillColor=c_title))

        cx, cy, r_out, r_in = 58, 146, 42, 27
        fatias = [
            ("Licenciamento", soma_lic, colors.HexColor("#F59E0B")),
            ("IPVA (Sefaz)", soma_ipva, colors.HexColor("#F43F5E")),
            ("Multas Detran", soma_multas, colors.HexColor("#EA580C")),
        ]
        positives = [(n, v, c) for n, v, c in fatias if v > 0]
        if soma_bruta <= 0 or not positives:
            d.add(Circle(cx, cy, r_out, fillColor=c_track, strokeColor=c_border, strokeWidth=1))
        elif len(positives) == 1:
            d.add(Circle(cx, cy, r_out, fillColor=positives[0][2], strokeColor=colors.white, strokeWidth=1))
        else:
            ang_atual = 90.0
            for _, val, cor in positives:
                ext = (val / soma_bruta) * 360.0
                d.add(Wedge(cx, cy, r_out, ang_atual - ext, ang_atual, fillColor=cor, strokeColor=colors.white, strokeWidth=1))
                ang_atual -= ext

        d.add(Circle(cx, cy, r_in, fillColor=c_bg_card, strokeColor=None))
        d.add(String(cx, cy + 4, "A Liquidar", fontName="Helvetica", fontSize=6, fillColor=c_muted, textAnchor="middle"))
        val_c_txt = formatar_float_moeda(soma_total).replace("R$ ", "")
        tam_c_f = 6.2 if len(val_c_txt) > 10 else 7.0
        d.add(String(cx, cy - 6, val_c_txt, fontName="Helvetica-Bold", fontSize=tam_c_f, fillColor=colors.HexColor("#15803D"), textAnchor="middle"))

        # Legenda lateral do Donut (x=110..242)
        ly = 172
        for nome, val, cor in fatias:
            pct = (val / soma_bruta * 100.0) if soma_bruta > 0 else 0.0
            d.add(Circle(114, ly + 2, 3.5, fillColor=cor, strokeColor=None))
            d.add(String(122, ly, f"{nome} ({pct:.0f}%)", fontName="Helvetica-Bold", fontSize=7.2, fillColor=c_title))
            d.add(String(122, ly - 10, formatar_float_moeda(val), fontName="Helvetica", fontSize=7.0, fillColor=c_sub))
            ly -= 30

        # Linha divisória e Resumo Financeiro inferior
        d.add(Line(12, 86, 238, 86, strokeColor=c_border, strokeWidth=0.6))
        ticket_dev = (soma_total / max(1, qtd_debitos)) if qtd_debitos > 0 else 0.0
        resumo_c1 = [
            ("Soma Bruta (Órgãos):", formatar_float_moeda(soma_bruta), c_sub),
            ("Abatimento Multas no Lic.:", f"- {formatar_float_moeda(abatimento)}", colors.HexColor("#15803D")),
            ("Total Líquido (Sem Duplic.):", formatar_float_moeda(soma_total), colors.HexColor("#DC2626" if soma_total > 0 else "#15803D")),
            ("Ticket Médio / Devedor:", formatar_float_moeda(ticket_dev), c_title),
        ]
        ry = 70
        for rot, val_s, cor_v in resumo_c1:
            d.add(String(12, ry, rot, fontName="Helvetica", fontSize=6.8, fillColor=c_sub))
            d.add(String(238, ry, val_s, fontName="Helvetica-Bold", fontSize=6.8, fillColor=cor_v, textAnchor="end"))
            ry -= 17

        # ======================================================================
        # CARD 2 (x=260..510): ÍNDICE DE REGULARIDADE DA FROTA
        # ======================================================================
        d.add(Rect(260, 0, 250, 226, rx=8, ry=8, fillColor=c_bg_card, strokeColor=c_border, strokeWidth=0.8))
        d.add(Rect(272, 223, 226, 3, fillColor=colors.HexColor("#10B981"), strokeColor=None))
        d.add(String(272, 207, "Índice de Regularidade da Frota", fontName="Helvetica-Bold", fontSize=8.5, fillColor=c_title))

        # Barra empilhada (x=272..498, w=226, y=184..195)
        bx1, bx2, bw = 272.0, 498.0, 226.0
        d.add(Rect(bx1, 184, bw, 11, fillColor=c_track, strokeColor=None))
        if auditados > 0:
            denom = max(total_veiculos, auditados, 1)
            x_cur = bx1
            for qtd, cor_hex in [(qtd_pagos, "#10B981"), (qtd_debitos, "#EF4444"), (qtd_falhas, "#F59E0B")]:
                if qtd > 0 and x_cur < bx2:
                    w_seg = min(bx2 - x_cur, (qtd / denom) * bw)
                    d.add(Rect(x_cur, 184, w_seg, 11, fillColor=colors.HexColor(cor_hex), strokeColor=None))
                    x_cur += w_seg

        pct_reg = (qtd_pagos / auditados * 100.0) if auditados > 0 else 0.0
        pct_deb = (qtd_debitos / auditados * 100.0) if auditados > 0 else 0.0

        # Grade 2x2 de Status
        grade_c2 = [
            (272, 164, "Regulares / Quitados", f"{qtd_pagos} veíc. ({pct_reg:.0f}%)", "#15803D"),
            (388, 164, "Com Débitos Ativos", f"{qtd_debitos} veíc. ({pct_deb:.0f}%)", "#DC2626"),
            (272, 126, "Com Restrição", f"{qtd_restricao} veíc.", "#EA580C"),
            (388, 126, "Falhas / Pendentes", f"{qtd_falhas} veíc.", "#D97706"),
        ]
        for gx, gy, rot, val_s, cor_h in grade_c2:
            d.add(String(gx, gy, rot, fontName="Helvetica", fontSize=7.0, fillColor=c_sub))
            d.add(String(gx, gy - 13, val_s, fontName="Helvetica-Bold", fontSize=8.8, fillColor=colors.HexColor(cor_h)))

        d.add(Line(272, 86, 498, 86, strokeColor=c_border, strokeWidth=0.6))
        custo_medio_frota = (soma_total / max(1, auditados)) if auditados > 0 else 0.0
        maior_passivo = f"{top10[0]['placa']} ({top10[0]['total_s']})" if top10 else "Nenhum"
        extras_c2 = [
            ("Veículos Auditados:", f"{auditados} de {total_veiculos}", c_title),
            ("Elegíveis p/ Emissão de Boleto:", f"{qtd_boletos} veíc.", colors.HexColor("#2563EB")),
            ("Custo Médio por Veículo:", formatar_float_moeda(custo_medio_frota), c_title),
            ("Maior Passivo Individual:", maior_passivo, colors.HexColor("#DC2626" if top10 else "#64748B")),
        ]
        ry = 70
        for rot, val_s, cor_v in extras_c2:
            d.add(String(272, ry, rot, fontName="Helvetica", fontSize=6.8, fillColor=c_sub))
            d.add(String(498, ry, val_s, fontName="Helvetica-Bold", fontSize=6.8, fillColor=cor_v, textAnchor="end"))
            ry -= 17

        # ======================================================================
        # CARD 3 (x=520..770): TOP 10 VEÍCULOS DE MAIOR IMPACTO
        # ======================================================================
        d.add(Rect(520, 0, 250, 226, rx=8, ry=8, fillColor=c_bg_card, strokeColor=c_border, strokeWidth=0.8))
        d.add(Rect(532, 223, 226, 3, fillColor=colors.HexColor("#3B82F6"), strokeColor=None))
        d.add(String(532, 207, "Top 10 Veículos de Maior Impacto", fontName="Helvetica-Bold", fontSize=8.5, fillColor=c_title))

        if not top10:
            d.add(String(645, 110, "Nenhum débito pendente registrado", fontName="Helvetica-Oblique", fontSize=8, fillColor=c_muted, textAnchor="middle"))
        else:
            maior_v = max((item["total_f"] for item in top10), default=1.0) or 1.0
            ry = 188.0
            passo_t = 18.0
            # Coluna 1: x=530..626 (máx 22 chars em Courier-Bold 6.4pt = 84pt)
            # Coluna 2 (Barra): x=630..700 (largura 70pt)
            # Coluna 3 (Valor): x=758 (anchor='end', ocupa ~708..758)
            bx_ini, larg_b = 630.0, 70.0
            for idx, item in enumerate(top10[:10]):
                rank = idx + 1
                placa = item["placa"]
                prop = str(item.get("prop") or "").strip()
                if prop and prop != "NÃO INFORMADO":
                    prop_curto = prop[:7] + "…" if len(prop) > 7 else prop
                    lbl_esq = f"{rank:>2}. {placa} ({prop_curto})"
                else:
                    lbl_esq = f"{rank:>2}. {placa}"

                d.add(String(530, ry, lbl_esq, fontName="Courier-Bold", fontSize=6.3, fillColor=c_title))

                frac = min(1.0, max(0.0, item["total_f"] / maior_v))
                w_fill = max(3.0, larg_b * frac)
                cor_b = "#EF4444" if rank == 1 else ("#F97316" if rank in (2, 3) else "#3B82F6")

                d.add(Rect(bx_ini, ry - 0.5, larg_b, 5.5, fillColor=c_track, strokeColor=None))
                d.add(Rect(bx_ini, ry - 0.5, w_fill, 5.5, fillColor=colors.HexColor(cor_b), strokeColor=None))

                d.add(String(758, ry, item["total_s"], fontName="Helvetica-Bold", fontSize=6.5, fillColor=c_sub, textAnchor="end"))
                ry -= passo_t

        return d
